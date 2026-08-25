"""Create the PDF project report and XLSX retrieval-evaluation workbook.

The script deliberately evaluates the existing question files against the
current corpus; it does not invent retrieval metrics.  It can be rerun after a
corpus or configuration change to regenerate both deliverables.
"""
from __future__ import annotations

import json
import os
import sys
from collections import defaultdict
from datetime import date
from html import escape
from pathlib import Path
from statistics import mean
from typing import Any

# The available multilingual model is already cached locally. Set this before
# importing SentenceTransformer through the retriever so report generation does
# not perform a metadata request to Hugging Face on restricted Windows networks.
os.environ.setdefault("HF_HUB_OFFLINE", "1")
os.environ.setdefault("TOKENIZERS_PARALLELISM", "false")

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from rag.chunking import make_chunks
from rag.evaluation import metrics
from rag.pdf_ingest import extract_pdf_pages
from rag.retrieval import SemanticRetriever


ROOT = PROJECT_ROOT
DATA_DIR = ROOT / "data"
PDF_DIR = DATA_DIR / "pdfs"
REPORT_PATH = ROOT / "Bao_cao_project_RAG_PDF.pdf"
WORKBOOK_PATH = ROOT / "Tap_cau_hoi_danh_gia_RAG.xlsx"
EMBEDDING_MODEL = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
CHUNK_SIZE = 700
CHUNK_OVERLAP = 120
TOP_K = 3

CATEGORY_LABELS = {
    "keyword_match": "Khớp từ khóa",
    "semantic_paraphrase": "Diễn đạt lại theo ngữ nghĩa",
    "multi_information": "Tổng hợp nhiều thông tin",
}
METRIC_COLUMNS = ("precision@3", "recall@3", "hit@3", "mrr", "ndcg@3", "query_ms")


def load_json(filename: str) -> list[dict[str, Any]]:
    return json.loads((DATA_DIR / filename).read_text(encoding="utf-8"))


def collect_corpus() -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """Return sections, file-level facts, and recursive chunks for the corpus."""
    sections: list[dict[str, Any]] = []
    file_rows: list[dict[str, Any]] = []
    for pdf_path in sorted(PDF_DIR.glob("*.pdf")):
        with pymupdf.open(pdf_path) as document:
            page_count = document.page_count
        document_sections = extract_pdf_pages(pdf_path)
        sections.extend(document_sections)
        file_rows.append(
            {
                "source": pdf_path.name,
                "pages": page_count,
                "sections": len(document_sections),
            }
        )
    chunks = make_chunks(sections, "recursive", CHUNK_SIZE, CHUNK_OVERLAP)
    return sections, file_rows, chunks


def ground_truth_text(ground_truth: list[dict[str, Any]], key: str, separator: str = "\n") -> str:
    return separator.join(str(item.get(key, "")) for item in ground_truth)


def page_range(chunk: dict[str, Any]) -> str:
    start = chunk.get("page_start", chunk.get("page", ""))
    end = chunk.get("page_end", start)
    return str(start) if start == end else f"{start}–{end}"


def run_evaluation(
    retriever: SemanticRetriever, items: list[dict[str, Any]], dataset_name: str, fallback_category: str | None = None
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Run Top-3 retrieval and retain both summary and inspectable hit rows."""
    rows: list[dict[str, Any]] = []
    hit_rows: list[dict[str, Any]] = []
    for item in items:
        hits, latency = retriever.search(item["question"], k=TOP_K)
        result = metrics([hit.chunk for hit in hits], item.get("ground_truth", []), k=TOP_K)
        category = item.get("category") or fallback_category or "other"
        ground_truth = item.get("ground_truth", [])
        top_one = hits[0].chunk if hits else {}
        row = {
            "dataset": dataset_name,
            "id": item.get("id", ""),
            "query_type": category,
            "query_type_label": CATEGORY_LABELS.get(category, category),
            "question": item["question"],
            "ground_truth_source": ground_truth_text(ground_truth, "source"),
            "ground_truth_page": ground_truth_text(ground_truth, "page"),
            "ground_truth_section": ground_truth_text(ground_truth, "section"),
            "answer_rubric": item.get("answer_rubric", ""),
            "config": f"recursive | size={CHUNK_SIZE} | overlap={CHUNK_OVERLAP} | top_k={TOP_K}",
            "query_ms": round(latency, 2),
            "top1_source": top_one.get("source", ""),
            "top1_pages": page_range(top_one) if top_one else "",
            "top1_section": top_one.get("section", ""),
        }
        row.update(result)
        rows.append(row)
        for hit in hits:
            chunk = hit.chunk
            hit_rows.append(
                {
                    "dataset": dataset_name,
                    "id": item.get("id", ""),
                    "query_type": CATEGORY_LABELS.get(category, category),
                    "question": item["question"],
                    "rank": hit.rank,
                    "similarity_score": round(float(hit.score), 4),
                    "chunk_id": chunk.get("chunk_id", ""),
                    "source": chunk.get("source", ""),
                    "page_start": chunk.get("page_start", chunk.get("page", "")),
                    "page_end": chunk.get("page_end", chunk.get("page", "")),
                    "section": chunk.get("section", ""),
                    "strategy": chunk.get("strategy", ""),
                    "text_preview": str(chunk.get("text", ""))[:700],
                }
            )
    return rows, hit_rows


def aggregate(rows: list[dict[str, Any]]) -> dict[str, float]:
    if not rows:
        return {metric: 0.0 for metric in METRIC_COLUMNS}
    return {metric: mean(float(row[metric]) for row in rows) for metric in METRIC_COLUMNS}


def aggregate_by_category(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row["query_type_label"]].append(row)
    summary: list[dict[str, Any]] = []
    for category, category_rows in grouped.items():
        data = {"query_type": category, "questions": len(category_rows)}
        data.update(aggregate(category_rows))
        summary.append(data)
    return summary


def safe_paragraph(value: Any) -> str:
    return escape(str(value)).replace("\n", "<br/>")


def register_fonts() -> tuple[str, str, str]:
    """Register Windows Arial fonts so Vietnamese content remains searchable."""
    font_directory = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    font_files = {
        "RAGViet": font_directory / "arial.ttf",
        "RAGVietBold": font_directory / "arialbd.ttf",
        "RAGVietItalic": font_directory / "ariali.ttf",
    }
    if all(path.exists() for path in font_files.values()):
        for name, path in font_files.items():
            pdfmetrics.registerFont(TTFont(name, str(path)))
        return "RAGViet", "RAGVietBold", "RAGVietItalic"
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"


def build_report_styles(font: str, bold: str, italic: str) -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "cover_title": ParagraphStyle(
            "CoverTitle", parent=base["Title"], fontName=bold, fontSize=24, leading=31, alignment=TA_CENTER,
            textColor=colors.HexColor("#12355B"), spaceAfter=18,
        ),
        "cover_subtitle": ParagraphStyle(
            "CoverSubtitle", parent=base["Normal"], fontName=font, fontSize=13, leading=19, alignment=TA_CENTER,
            textColor=colors.HexColor("#355C7D"),
        ),
        "heading": ParagraphStyle(
            "Heading", parent=base["Heading1"], fontName=bold, fontSize=16, leading=21, textColor=colors.HexColor("#12355B"),
            spaceAfter=10,
        ),
        "subheading": ParagraphStyle(
            "Subheading", parent=base["Heading2"], fontName=bold, fontSize=11.5, leading=15, textColor=colors.HexColor("#1F5C8D"),
            spaceBefore=6, spaceAfter=5,
        ),
        "body": ParagraphStyle(
            "Body", parent=base["BodyText"], fontName=font, fontSize=9.4, leading=14, alignment=TA_LEFT, spaceAfter=6,
        ),
        "bullet": ParagraphStyle(
            "Bullet", parent=base["BodyText"], fontName=font, fontSize=9.2, leading=13.2, leftIndent=12, firstLineIndent=-8,
            spaceAfter=3,
        ),
        "table_header": ParagraphStyle(
            "TableHeader", parent=base["Normal"], fontName=bold, fontSize=7.6, leading=9.3, textColor=colors.white,
        ),
        "table_cell": ParagraphStyle(
            "TableCell", parent=base["Normal"], fontName=font, fontSize=7.4, leading=9.2,
        ),
        "metric": ParagraphStyle(
            "Metric", parent=base["Normal"], fontName=bold, fontSize=12, leading=15, alignment=TA_CENTER,
            textColor=colors.HexColor("#12355B"),
        ),
        "small": ParagraphStyle(
            "Small", parent=base["Normal"], fontName=italic, fontSize=7.5, leading=10, textColor=colors.HexColor("#566573"),
        ),
        "flow": ParagraphStyle(
            "Flow", parent=base["Normal"], fontName=bold, fontSize=8.6, leading=11, alignment=TA_CENTER,
            textColor=colors.HexColor("#12355B"),
        ),
    }


def paragraph(text: Any, styles: dict[str, ParagraphStyle], style: str = "body") -> Paragraph:
    return Paragraph(safe_paragraph(text), styles[style])


def report_table(
    rows: list[list[Any]], widths: list[float], styles: dict[str, ParagraphStyle], *, header: bool = True
) -> Table:
    flow_rows: list[list[Paragraph]] = []
    for row_index, row in enumerate(rows):
        cell_style = "table_header" if header and row_index == 0 else "table_cell"
        flow_rows.append([paragraph(value, styles, cell_style) for value in row])
    table = Table(flow_rows, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    commands: list[tuple[Any, ...]] = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B8C7D9")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if header:
        commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F5C8D")),
                ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F7FAFC")),
            ]
        )
    table.setStyle(TableStyle(commands))
    return table


def metric_value(value: float, digits: int = 3) -> str:
    return f"{value:.{digits}f}"


def report_footer(canvas: Any, doc: Any, font: str) -> None:
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#B8C7D9"))
    canvas.setLineWidth(0.4)
    canvas.line(1.7 * cm, 1.35 * cm, A4[0] - 1.7 * cm, 1.35 * cm)
    canvas.setFont(font, 7.2)
    canvas.setFillColor(colors.HexColor("#566573"))
    canvas.drawString(1.7 * cm, 0.92 * cm, "Báo cáo Project RAG Semantic Search trên PDF")
    canvas.drawRightString(A4[0] - 1.7 * cm, 0.92 * cm, f"Trang {doc.page}")
    canvas.restoreState()


def build_pdf_report(
    file_rows: list[dict[str, Any]],
    sections: list[dict[str, Any]],
    chunks: list[dict[str, Any]],
    benchmark_rows: list[dict[str, Any]],
    ui_rows: list[dict[str, Any]],
    analysis_cases: list[dict[str, Any]],
) -> int:
    font, bold, italic = register_fonts()
    styles = build_report_styles(font, bold, italic)
    document = SimpleDocTemplate(
        str(REPORT_PATH),
        pagesize=A4,
        leftMargin=1.7 * cm,
        rightMargin=1.7 * cm,
        topMargin=1.55 * cm,
        bottomMargin=1.65 * cm,
        title="Báo cáo Project RAG Semantic Search trên PDF",
        author="Nhóm thực hiện Project CĐCNTT",
    )
    story: list[Any] = []
    corpus_pages = sum(int(item["pages"]) for item in file_rows)
    fixed_chunks = make_chunks(sections, "fixed", CHUNK_SIZE, CHUNK_OVERLAP)
    recursive_average = mean(len(chunk["text"]) for chunk in chunks)
    fixed_average = mean(len(chunk["text"]) for chunk in fixed_chunks)
    benchmark_average = aggregate(benchmark_rows)
    ui_average = aggregate(ui_rows)
    category_summary = aggregate_by_category(benchmark_rows)

    # Page 1 — cover.
    story.extend(
        [
            Spacer(1, 3.4 * cm),
            paragraph("BÁO CÁO PROJECT", styles, "cover_subtitle"),
            Spacer(1, 0.35 * cm),
            paragraph("HỆ THỐNG HỎI ĐÁP PDF\nDỰA TRÊN SEMANTIC SEARCH VÀ RAG", styles, "cover_title"),
            Spacer(1, 0.35 * cm),
            paragraph("SentenceTransformer · FAISS · Chunking theo section · Citation theo trang", styles, "cover_subtitle"),
            Spacer(1, 1.35 * cm),
            report_table(
                [
                    ["Nội dung", "Thông tin"],
                    ["Phạm vi", "Tra cứu song ngữ trên PDF và trả lời có căn cứ nguồn"],
                    ["Corpus", f"{len(file_rows)} PDF · {corpus_pages} trang · {len(sections)} sections"],
                    ["Cấu hình thực nghiệm", f"Recursive · {CHUNK_SIZE} ký tự · overlap {CHUNK_OVERLAP} · Top-{TOP_K}"],
                    ["Ngày tạo báo cáo", date.today().isoformat()],
                ],
                [4.4 * cm, 11.5 * cm],
                styles,
            ),
            Spacer(1, 1.25 * cm),
            paragraph("Báo cáo được sinh từ source code, corpus PDF và các bộ câu hỏi hiện có trong project.", styles, "small"),
            PageBreak(),
        ]
    )

    # Page 2 — problem.
    story.extend(
        [
            paragraph("1. Bài toán cụ thể của project", styles, "heading"),
            paragraph(
                "Project giải quyết bài toán hỏi đáp dựa trên bộ tài liệu PDF tiếng Việt và tiếng Anh. Người dùng nhập một câu hỏi; hệ thống phải tìm được bằng chứng liên quan trong tài liệu, sinh câu trả lời bám sát bằng chứng và đưa lại tên file, mục/chương, cùng số trang để người dùng kiểm tra.",
                styles,
            ),
            paragraph("Vấn đề cần khắc phục", styles, "subheading"),
            *[
                paragraph(item, styles, "bullet")
                for item in [
                    "• PDF có ngắt trang, xuống dòng, khoảng trắng dư thừa và từ bị ngắt dòng; nếu nối toàn bộ file sẽ mất provenance theo trang.",
                    "• Các đề cương học phần có các section rất giống nhau như “Mục tiêu học phần”; truy hồi toàn corpus dễ lấy nhầm học phần/tài liệu.",
                    "• Bảng PDF và các heading in đậm/in nghiêng hoặc “Ghi chú/Lưu ý” cần trở thành evidence có thể truy hồi, thay vì bị bỏ qua.",
                    "• Câu trả lời phải trực tiếp, có citation đúng context, không dựa vào kiến thức bên ngoài tài liệu.",
                ]
            ],
            paragraph("Mục tiêu đầu ra", styles, "subheading"),
            report_table(
                [
                    ["Mục tiêu", "Cách đáp ứng trong project"],
                    ["Tìm kiếm ngữ nghĩa", "Embedding đa ngôn ngữ + FAISS cosine similarity; không sử dụng BM25/RRF."],
                    ["Căn cứ nguồn", "Mỗi chunk/section giữ source, page_start, page_end, pages, section và section_id."],
                    ["Hỏi đa dạng", "Hỗ trợ định nghĩa, giải thích, ví dụ, so sánh, tra theo mục/trang và câu hỏi nhiều thành phần."],
                    ["Kiểm chứng", "Hiển thị Top-K, nội dung đầy đủ của section, metrics retrieval và thời gian truy vấn."],
                ],
                [4.0 * cm, 11.9 * cm],
                styles,
            ),
            PageBreak(),
        ]
    )

    # Page 3 — data.
    story.extend(
        [
            paragraph("2. Dữ liệu của project", styles, "heading"),
            paragraph(
                f"Corpus hiện có {len(file_rows)} file PDF với tổng {corpus_pages} trang. Hệ thống kiểm tra tối thiểu 3 PDF và mỗi PDF từ 5 trang trước khi lập chỉ mục.",
                styles,
            ),
            report_table(
                [["Tài liệu", "Số trang", "Sections sau trích xuất"]]
                + [[row["source"], row["pages"], row["sections"]] for row in file_rows],
                [9.4 * cm, 2.7 * cm, 3.8 * cm],
                styles,
            ),
            Spacer(1, 0.15 * cm),
            paragraph("Tập đánh giá và phân tích", styles, "subheading"),
            report_table(
                [
                    ["File", "Vai trò", "Quy mô"],
                    ["evaluation_questions.json", "Ground truth đang dùng trên giao diện; các câu nhiều thành phần.", "3 câu / 7 sections GT"],
                    ["evaluation_questions_15.json", "Benchmark cân bằng theo loại query.", "15 câu: 5/5/5"],
                    ["retrieval_analysis_cases.json", "Phân tích case tốt, case mơ hồ và khác biệt chiến lược.", "3 cases"],
                ],
                [4.7 * cm, 7.8 * cm, 3.4 * cm],
                styles,
            ),
            paragraph("Tính chất dữ liệu", styles, "subheading"),
            *[
                paragraph(item, styles, "bullet")
                for item in [
                    "• Có tài liệu tiếng Việt, tiếng Anh và tài liệu hỗn hợp; đặc biệt có đề cương học phần với cấu trúc lặp lại.",
                    "• Có bảng đánh giá, danh mục tài liệu, danh sách cán bộ giảng dạy và các section trải qua nhiều trang.",
                    "• Ground truth lưu theo source/page/section để người đánh giá có thể mở PDF kiểm tra trực tiếp.",
                ]
            ],
            PageBreak(),
        ]
    )

    # Page 4 — architecture.
    flow_rows = []
    for index, label in enumerate(
        [
            "PDF theo từng trang",
            "Làm sạch + bảng + heading",
            "Ghép section qua page break",
            "Chunk + metadata",
            "Embedding đa ngôn ngữ",
            "FAISS Semantic Search",
            "Top-K + deduplicate section",
            "RAG / citation / UI",
        ]
    ):
        flow_rows.append([paragraph(label, styles, "flow")])
        if index < 7:
            flow_rows.append([paragraph("↓", styles, "flow")])
    flow = Table(flow_rows, colWidths=[15.9 * cm], hAlign="CENTER")
    flow.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#8DA9C4")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF2F8")),
                ("BACKGROUND", (0, 2), (-1, 2), colors.HexColor("#F7FAFC")),
                ("BACKGROUND", (0, 4), (-1, 4), colors.HexColor("#EAF2F8")),
                ("BACKGROUND", (0, 6), (-1, 6), colors.HexColor("#F7FAFC")),
                ("BACKGROUND", (0, 8), (-1, 8), colors.HexColor("#EAF2F8")),
                ("BACKGROUND", (0, 10), (-1, 10), colors.HexColor("#F7FAFC")),
                ("BACKGROUND", (0, 12), (-1, 12), colors.HexColor("#EAF2F8")),
                ("BACKGROUND", (0, 14), (-1, 14), colors.HexColor("#F7FAFC")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend(
        [
            paragraph("3. Thiết kế hệ thống", styles, "heading"),
            paragraph("Hệ thống được thiết kế theo pipeline page-aware RAG: metadata không bị mất trên đường từ PDF đến UI và LLM.", styles),
            flow,
            Spacer(1, 0.16 * cm),
            report_table(
                [
                    ["Thành phần", "Vai trò", "File chính"],
                    ["PDF ingestion", "Trích từng trang, table extraction, heading detection, section reconstruction.", "rag/pdf_ingest.py"],
                    ["Chunking", "Fixed-size hoặc recursive, kế thừa metadata của section.", "rag/chunking.py"],
                    ["Retrieval", "SentenceTransformer, FAISS IndexFlatIP, scope nguồn và deduplicate section.", "rag/retrieval.py"],
                    ["Generation", "Context-only prompt, citation từ metadata retrieved.", "rag/generation.py"],
                    ["Evaluation & UI", "Metrics section-level, Top-K debug, thống kê chunk/latency.", "rag/evaluation.py / app.py"],
                ],
                [3.4 * cm, 8.3 * cm, 4.2 * cm],
                styles,
            ),
            PageBreak(),
        ]
    )

    # Page 5 — parsing and chunking.
    chunk_example = '{"chunk_id": 17, "text": "…", "source": "document.pdf", "page": 5,\n "page_start": 5, "page_end": 6, "section": "3. Mục tiêu học phần",\n "section_id": "…", "strategy": "recursive"}'
    story.extend(
        [
            paragraph("4. Xử lý PDF, bảng và chunking", styles, "heading"),
            paragraph(
                "PDF không được nối thành một chuỗi duy nhất. Hệ thống đọc từng trang, giữ thứ tự đọc theo toạ độ, sau đó chỉ kết thúc section khi gặp heading mới hoặc hết tài liệu. Vì vậy nội dung section kéo dài qua page break vẫn có page_start/page_end chính xác.",
                styles,
            ),
            report_table(
                [
                    ["Bước", "Xử lý", "Lợi ích"],
                    ["Làm sạch", "Whitespace/dòng trống dư thừa; ghép hyphenation; khôi phục khoảng cách glyph.", "Đoạn text đọc tự nhiên hơn."],
                    ["Heading", "Nhận dạng số thứ tự, chương, heading bold/italic, Ghi chú/Lưu ý/Note.", "Không nhập nhầm nội dung mục trước."],
                    ["Bảng", "Chuyển thành các cặp header: value và giữ cell để UI render DataFrame.", "Truy hồi được trọng số, tiêu chí và nội dung bảng."],
                    ["Section-aware", "Ghép section qua trang; recursive chỉ tách trong một section lớn.", "Giữ mạch ngữ nghĩa và citation."],
                ],
                [2.5 * cm, 8.0 * cm, 5.4 * cm],
                styles,
            ),
            paragraph("Metadata của mỗi chunk", styles, "subheading"),
            report_table([[chunk_example]], [15.9 * cm], styles, header=False),
            Spacer(1, 0.16 * cm),
            report_table(
                [
                    ["Chiến lược", "Số chunk", "Độ dài TB (ký tự)", "Nhận xét"],
                    ["Recursive", len(chunks), f"{recursive_average:.1f}", "Ưu tiên ranh giới câu/section."],
                    ["Fixed-size", len(fixed_chunks), f"{fixed_average:.1f}", "Cắt theo ký tự, có overlap."],
                ],
                [3.2 * cm, 3.0 * cm, 3.8 * cm, 5.9 * cm],
                styles,
            ),
            PageBreak(),
        ]
    )

    # Page 6 — retrieval, RAG, UI.
    story.extend(
        [
            paragraph("5. Retrieval, RAG và giao diện", styles, "heading"),
            paragraph(
                "Embedding sử dụng model sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2. Vector được normalize và đưa vào FAISS IndexFlatIP, tương đương cosine similarity. Text được embed gồm section và nội dung chunk để tiêu đề mục tham gia vào độ tương đồng.",
                styles,
            ),
            report_table(
                [
                    ["Khâu", "Thiết kế hiện có"],
                    ["Nhận diện nguồn", "Tạo alias từ filename; nếu phát hiện tên file/học phần, candidate bị giới hạn đúng source. UI vẫn cho phép ưu tiên thủ công."],
                    ["Truy hồi", "Sinh embedding query, tìm FAISS, kết hợp điểm embedding nội dung/section, ưu tiên ý định section và loại trùng section."],
                    ["Context", "Top-K unique sections được ghép không trùng overlap; hiển thị cả chunk và full cleaned section."],
                    ["LLM", "Gemini (GEMINI_API_KEY, mặc định gemini-2.5-flash) chỉ dùng context; thiếu key vẫn xem được Top-K."],
                    ["Citation", "Câu trả lời mang citation [PDF | trang | mục] lấy từ metadata context, không để LLM tự tạo source."],
                ],
                [3.2 * cm, 12.7 * cm],
                styles,
            ),
            paragraph("Khả năng kiểm tra trên UI", styles, "subheading"),
            *[
                paragraph(item, styles, "bullet")
                for item in [
                    "• Nạp/lập chỉ mục, chọn recursive hoặc fixed-size, điều chỉnh size/overlap và Top-K.",
                    "• Hiển thị similarity, chunk_id, source, section, page range, strategy, text và bảng PDF gốc.",
                    "• Hiển thị latency, ước lượng context token, số chunk, độ dài trung bình, Precision@3 và Recall@3 khi câu hỏi có ground truth.",
                ]
            ],
            PageBreak(),
        ]
    )

    # Page 7 — experiment design.
    story.extend(
        [
            paragraph("6. Thực nghiệm và phương pháp đánh giá", styles, "heading"),
            paragraph(
                "Thực nghiệm retrieval chạy trên toàn bộ corpus với Top-3 và dùng ground truth theo section. Một result được xem là đúng khi trùng cặp (source, section); page được lưu để kiểm tra tài liệu nhưng không tham gia so khớp metric.",
                styles,
            ),
            report_table(
                [
                    ["Tham số", "Giá trị"],
                    ["Embedding model", EMBEDDING_MODEL],
                    ["Index", "FAISS IndexFlatIP + normalized vectors (cosine similarity)"],
                    ["Chunk chính", f"recursive, size={CHUNK_SIZE}, overlap={CHUNK_OVERLAP}"],
                    ["Final Top-K", TOP_K],
                    ["Benchmark chính", "evaluation_questions_15.json: 5 keyword + 5 semantic + 5 multi-information"],
                ],
                [5.0 * cm, 10.9 * cm],
                styles,
            ),
            paragraph("Các metric", styles, "subheading"),
            report_table(
                [
                    ["Metric", "Ý nghĩa"],
                    ["Precision@3", "Tỷ lệ section đúng trong ba kết quả đầu; mẫu số cố định là 3."],
                    ["Recall@3", "Tỷ lệ toàn bộ section ground truth được lấy được trong Top-3."],
                    ["Hit@3 / MRR", "Có hit hay không; và thứ hạng của hit đúng đầu tiên."],
                    ["nDCG@3", "Đánh giá chất lượng thứ hạng khi có nhiều section đúng."],
                    ["Query latency", "Thời gian semantic retrieval, không gồm thời gian LLM sinh câu trả lời."],
                ],
                [3.4 * cm, 12.5 * cm],
                styles,
            ),
            paragraph(
                "Ngoài benchmark định lượng, project lưu một case tốt, một case truy vấn mơ hồ không có ground truth duy nhất, và một case so sánh recursive/fixed. Điều này tránh kết luận chunking chỉ dựa vào cảm quan.",
                styles,
            ),
            PageBreak(),
        ]
    )

    # Page 8 — numeric results.
    story.extend(
        [
            paragraph("7. Kết quả thực nghiệm", styles, "heading"),
            paragraph(
                "Bảng dưới đây là kết quả chạy lại khi sinh báo cáo trên corpus hiện có, với model cache cục bộ và cấu hình recursive 700/120. Chi tiết từng câu, ground truth và Top-3 được xuất trong workbook XLSX đi kèm.",
                styles,
            ),
            report_table(
                [
                    ["Bộ câu hỏi", "Số câu", "P@3", "R@3", "Hit@3", "MRR", "nDCG@3", "Latency TB (ms)"],
                    [
                        "Benchmark 15 câu",
                        len(benchmark_rows),
                        metric_value(benchmark_average["precision@3"]),
                        metric_value(benchmark_average["recall@3"]),
                        metric_value(benchmark_average["hit@3"]),
                        metric_value(benchmark_average["mrr"]),
                        metric_value(benchmark_average["ndcg@3"]),
                        metric_value(benchmark_average["query_ms"], 1),
                    ],
                    [
                        "3 câu dùng trên UI",
                        len(ui_rows),
                        metric_value(ui_average["precision@3"]),
                        metric_value(ui_average["recall@3"]),
                        metric_value(ui_average["hit@3"]),
                        metric_value(ui_average["mrr"]),
                        metric_value(ui_average["ndcg@3"]),
                        metric_value(ui_average["query_ms"], 1),
                    ],
                ],
                [3.0 * cm, 1.35 * cm, 1.45 * cm, 1.45 * cm, 1.45 * cm, 1.35 * cm, 1.55 * cm, 2.3 * cm],
                styles,
            ),
            Spacer(1, 0.2 * cm),
            paragraph("Phân rã benchmark 15 câu theo loại query", styles, "subheading"),
            report_table(
                [["Loại query", "Số câu", "P@3", "R@3", "Hit@3", "MRR", "nDCG@3", "Latency TB (ms)"]]
                + [
                    [
                        row["query_type"], row["questions"], metric_value(row["precision@3"]),
                        metric_value(row["recall@3"]), metric_value(row["hit@3"]), metric_value(row["mrr"]),
                        metric_value(row["ndcg@3"]), metric_value(row["query_ms"], 1),
                    ]
                    for row in category_summary
                ],
                [3.4 * cm, 1.2 * cm, 1.4 * cm, 1.4 * cm, 1.4 * cm, 1.3 * cm, 1.5 * cm, 2.3 * cm],
                styles,
            ),
            Spacer(1, 0.16 * cm),
            paragraph(
                "Kết quả cần được đọc cùng section/source của Top-K. Precision/Recall phản ánh evidence retrieval; chúng không thay thế việc đánh giá đúng–đủ–bám nguồn của câu trả lời LLM.",
                styles,
            ),
            PageBreak(),
        ]
    )

    # Page 9 — analysis cases.
    case_by_type = {case.get("case_type"): case for case in analysis_cases}
    good_case = case_by_type.get("good", {})
    failure_case = case_by_type.get("failure", {})
    strategy_case = case_by_type.get("strategy_difference", {})
    story.extend(
        [
            paragraph("8. Phân tích ba tình huống tiêu biểu", styles, "heading"),
            paragraph("Case tốt", styles, "subheading"),
            paragraph(
                f"Câu hỏi: “{good_case.get('question', '')}”. Tên học phần trong query cho phép nhận diện source, còn nội dung câu hỏi khớp section tài liệu tham khảo. Tiêu chí thành công là Top-1 thuộc đúng section và câu trả lời nêu đúng giáo trình kèm citation.",
                styles,
            ),
            paragraph("Case thất bại / mơ hồ", styles, "subheading"),
            paragraph(
                f"Câu hỏi: “{failure_case.get('question', '')}”. Nhiều đề cương có section “Mục tiêu học phần”, trong khi query không nêu file hoặc học phần. Vì không có scope metadata, không nên gán một ground truth đơn lẻ hay tính Precision/Recall cho case này; hành vi mong muốn là yêu cầu người dùng làm rõ.",
                styles,
            ),
            paragraph("Case recursive và fixed-size khác nhau", styles, "subheading"),
            paragraph(
                f"Câu hỏi: “{strategy_case.get('question', '')}”. Dữ liệu case lưu trong project ghi nhận recursive đưa Top-1 bắt đầu trực tiếp tại heading và nội dung phát minh Công nghiệp 4.0 (score 0.8172). Fixed-size có score cao hơn (0.8469) nhưng Top-1 mở đầu bằng phần cuối không liên quan rồi mới đến heading. Điều này cho thấy score đơn lẻ chưa đủ: tính hoàn chỉnh của evidence ảnh hưởng trực tiếp chất lượng câu trả lời.",
                styles,
            ),
            report_table(
                [
                    ["Tiêu chí", "Recursive", "Fixed-size"],
                    ["Ranh giới", "Tôn trọng câu/section khi có thể.", "Cắt theo kích thước ký tự."],
                    ["Ngữ cảnh Top-1", "Bắt đầu trực tiếp bằng heading liên quan.", "Có thể kèm phần cuối section trước."],
                    ["Kết luận", "Phù hợp hơn cho citation/answer coherence trong case dài.", "Hữu ích baseline, nhưng cần kiểm tra overlap và boundary."],
                ],
                [3.2 * cm, 6.35 * cm, 6.35 * cm],
                styles,
            ),
            Spacer(1, 0.16 * cm),
            paragraph("Vì vậy chiến lược chunking được so sánh bằng retrieval metrics, Top-K evidence và chất lượng context, không chỉ bằng số similarity hay quan sát trực quan.", styles),
            PageBreak(),
        ]
    )

    # Page 10 — limitations and conclusion.
    story.extend(
        [
            paragraph("9. Hạn chế và kết luận", styles, "heading"),
            paragraph("Hạn chế", styles, "subheading"),
            *[
                paragraph(item, styles, "bullet")
                for item in [
                    "• Ground truth hiện tập trung nhiều vào đề cương Tiếng Anh thương mại 2; cần mở rộng sang các PDF và loại nội dung khác.",
                    "• Query quá mơ hồ, không nêu file/học phần, vẫn có thể trả về nhiều section tương tự từ các tài liệu khác.",
                    "• PDF scan hoặc bảng có layout phức tạp có thể cần OCR/layout parser chuyên dụng hơn PyMuPDF.",
                    "• Project không dùng reranker; chất lượng xếp hạng phụ thuộc embedding, metadata và section-aware scoring.",
                    "• Metrics hiện đo retrieval evidence, chưa tự động chấm đầy đủ tính đúng/đủ/citation của câu trả lời LLM.",
                ]
            ],
            paragraph("Kết luận", styles, "subheading"),
            paragraph(
                "Project đã xây dựng một hệ thống RAG semantic-search có khả năng đọc PDF theo trang, phục hồi section qua page break, truy hồi bảng và heading, giữ metadata đầy đủ, và trả lời có trích dẫn kiểm tra được. Thiết kế section-aware cùng recursive chunking làm evidence liền mạch hơn, đặc biệt đối với các đề cương có cấu trúc lặp lại.",
                styles,
            ),
            paragraph("Hướng phát triển", styles, "subheading"),
            paragraph(
                "Mở rộng tập ground truth đa tài liệu; bổ sung OCR cho scan, detector bảng/layout mạnh hơn, reranker semantic chuyên dụng, và đánh giá answer-level bằng rubric có người kiểm tra. Khi thêm PDF mới, cần tái trích xuất, lập chỉ mục và cập nhật câu hỏi ground truth tương ứng.",
                styles,
            ),
            Spacer(1, 0.2 * cm),
            report_table(
                [
                    ["Tệp bàn giao", "Nội dung"],
                    [REPORT_PATH.name, "Báo cáo 10 trang về bài toán, dữ liệu, thiết kế, thực nghiệm, kết quả, hạn chế và kết luận."],
                    [WORKBOOK_PATH.name, "Tập 15 câu + 3 câu UI, ground truth, Top-3 và metrics truy hồi."],
                ],
                [5.1 * cm, 10.8 * cm],
                styles,
            ),
        ]
    )

    document.build(
        story,
        onFirstPage=lambda canvas, doc: report_footer(canvas, doc, font),
        onLaterPages=lambda canvas, doc: report_footer(canvas, doc, font),
    )
    with pymupdf.open(REPORT_PATH) as generated:
        page_count = generated.page_count
    if not 8 <= page_count <= 15:
        raise RuntimeError(f"PDF report must be 8–15 pages, generated {page_count} pages.")
    return page_count


def set_sheet_layout(sheet: Any, *, landscape: bool = True) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.45
    sheet.page_margins.bottom = 0.45


def append_table(sheet: Any, headers: list[str], rows: list[list[Any]], *, start_row: int = 1) -> int:
    header_fill = PatternFill("solid", fgColor="1F5C8D")
    even_fill = PatternFill("solid", fgColor="F3F7FB")
    thin = Side(style="thin", color="B8C7D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(start_row, column, value)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row_number, values in enumerate(rows, start=start_row + 1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if (row_number - start_row) % 2 == 0:
                cell.fill = even_fill
        sheet.row_dimensions[row_number].height = 42
    sheet.row_dimensions[start_row].height = 32
    for column in range(1, len(headers) + 1):
        values = [str(sheet.cell(row, column).value or "") for row in range(start_row, start_row + len(rows) + 1)]
        width = min(max(max(len(value) for value in values) + 2, 11), 55)
        sheet.column_dimensions[get_column_letter(column)].width = width
    return start_row + len(rows)


def write_workbook(
    file_rows: list[dict[str, Any]],
    benchmark_rows: list[dict[str, Any]],
    ui_rows: list[dict[str, Any]],
    hit_rows: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Tong_hop"
    benchmark_average = aggregate(benchmark_rows)
    ui_average = aggregate(ui_rows)
    summary_rows = [
        ["Ngày tạo", date.today().isoformat()],
        ["Corpus", f"{len(file_rows)} PDF / {sum(row['pages'] for row in file_rows)} trang"],
        ["Embedding model", EMBEDDING_MODEL],
        ["Cấu hình", f"recursive | chunk size {CHUNK_SIZE} | overlap {CHUNK_OVERLAP} | Top-{TOP_K}"],
        ["Benchmark chính", "data/evaluation_questions_15.json (15 câu)"],
        ["Bộ UI", "data/evaluation_questions.json (3 câu)"],
    ]
    summary_end = append_table(summary_sheet, ["Thông tin", "Giá trị"], summary_rows)
    summary_sheet.column_dimensions["A"].width = 28
    summary_sheet.column_dimensions["B"].width = 90
    summary_end += 3
    summary_metric_rows = [
        ["Benchmark 15 câu", len(benchmark_rows)] + [round(benchmark_average[metric], 4) for metric in METRIC_COLUMNS],
        ["3 câu UI", len(ui_rows)] + [round(ui_average[metric], 4) for metric in METRIC_COLUMNS],
    ]
    summary_end = append_table(
        summary_sheet,
        ["Tập", "Số câu", "Precision@3", "Recall@3", "Hit@3", "MRR", "nDCG@3", "Latency TB (ms)"],
        summary_metric_rows,
        start_row=summary_end,
    )
    summary_end += 3
    append_table(
        summary_sheet,
        ["Loại query", "Số câu", "Precision@3", "Recall@3", "Hit@3", "MRR", "nDCG@3", "Latency TB (ms)"],
        [
            [row["query_type"], row["questions"]] + [round(row[metric], 4) for metric in METRIC_COLUMNS]
            for row in aggregate_by_category(benchmark_rows)
        ],
        start_row=summary_end,
    )
    set_sheet_layout(summary_sheet)

    question_headers = [
        "Nguồn tập", "ID", "Loại query", "Câu hỏi", "GT source", "GT trang", "GT section", "Answer rubric",
        "Cấu hình", "Precision@3", "Recall@3", "Hit@3", "MRR", "nDCG@3", "Latency (ms)",
        "GT đúng lấy được", "Tổng GT section", "Top-1 source", "Top-1 trang", "Top-1 section",
    ]

    def question_values(rows: list[dict[str, Any]]) -> list[list[Any]]:
        return [
            [
                row["dataset"], row["id"], row["query_type_label"], row["question"], row["ground_truth_source"],
                row["ground_truth_page"], row["ground_truth_section"], row["answer_rubric"], row["config"],
                round(row["precision@3"], 4), round(row["recall@3"], 4), round(row["hit@3"], 4),
                round(row["mrr"], 4), round(row["ndcg@3"], 4), row["query_ms"], row["relevant_retrieved"],
                row["relevant_ground_truth"], row["top1_source"], row["top1_pages"], row["top1_section"],
            ]
            for row in rows
        ]

    benchmark_sheet = workbook.create_sheet("Benchmark_15")
    append_table(benchmark_sheet, question_headers, question_values(benchmark_rows))
    set_sheet_layout(benchmark_sheet)
    benchmark_sheet.freeze_panes = "D2"

    ui_sheet = workbook.create_sheet("UI_3_cau")
    append_table(ui_sheet, question_headers, question_values(ui_rows))
    set_sheet_layout(ui_sheet)
    ui_sheet.freeze_panes = "D2"

    hit_sheet = workbook.create_sheet("Top3_chi_tiet")
    hit_headers = [
        "Nguồn tập", "Question ID", "Loại query", "Câu hỏi", "Rank", "Similarity", "Chunk ID", "Source",
        "Trang bắt đầu", "Trang kết thúc", "Section", "Strategy", "Text preview",
    ]
    append_table(
        hit_sheet,
        hit_headers,
        [
            [
                row["dataset"], row["id"], row["query_type"], row["question"], row["rank"],
                row["similarity_score"], row["chunk_id"], row["source"], row["page_start"], row["page_end"],
                row["section"], row["strategy"], row["text_preview"],
            ]
            for row in hit_rows
        ],
    )
    set_sheet_layout(hit_sheet)
    hit_sheet.freeze_panes = "D2"

    guide_sheet = workbook.create_sheet("Huong_dan")
    guide_rows = [
        ["Mục đích", "Workbook lưu bộ câu hỏi đánh giá, ground truth và metrics retrieval được chạy trực tiếp trên corpus."],
        ["Đơn vị đúng", "Section: một hit đúng khi trùng source + section với ground truth. Page dùng để kiểm tra PDF."],
        ["Precision@3", "Số section đúng trong Top-3 chia cho 3."],
        ["Recall@3", "Số section GT đã lấy được trong Top-3 chia cho tổng số section GT."],
        ["Hit@3", "1 nếu có ít nhất một section đúng trong Top-3; ngược lại 0."],
        ["MRR", "Nghịch đảo thứ hạng của hit đúng đầu tiên."],
        ["nDCG@3", "Chất lượng thứ hạng khi truy vấn có nhiều section ground truth."],
        ["Latency", "Thời gian retrieval semantic, đơn vị ms; không gồm LLM generation."],
        ["Cấu hình", f"Model {EMBEDDING_MODEL}; recursive; size={CHUNK_SIZE}; overlap={CHUNK_OVERLAP}; Top-{TOP_K}."],
        ["Nguồn", "Benchmark_15 lấy từ evaluation_questions_15.json; UI_3_cau lấy từ evaluation_questions.json."],
    ]
    append_table(guide_sheet, ["Trường", "Giải thích"], guide_rows)
    guide_sheet.column_dimensions["A"].width = 24
    guide_sheet.column_dimensions["B"].width = 105
    set_sheet_layout(guide_sheet, landscape=False)

    workbook.save(WORKBOOK_PATH)


def main() -> None:
    benchmark_questions = load_json("evaluation_questions_15.json")
    ui_questions = load_json("evaluation_questions.json")
    analysis_cases = load_json("retrieval_analysis_cases.json")
    sections, file_rows, chunks = collect_corpus()
    retriever = SemanticRetriever(chunks, sections=sections, model_name=EMBEDDING_MODEL)
    benchmark_rows, benchmark_hits = run_evaluation(retriever, benchmark_questions, "evaluation_questions_15.json")
    ui_rows, ui_hits = run_evaluation(
        retriever, ui_questions, "evaluation_questions.json", fallback_category="multi_information"
    )
    write_workbook(file_rows, benchmark_rows, ui_rows, benchmark_hits + ui_hits)
    report_pages = build_pdf_report(file_rows, sections, chunks, benchmark_rows, ui_rows, analysis_cases)
    print(f"Created: {REPORT_PATH.name} ({report_pages} pages)")
    print(f"Created: {WORKBOOK_PATH.name} ({len(benchmark_rows)} benchmark questions + {len(ui_rows)} UI questions)")


if __name__ == "__main__":
    main()
